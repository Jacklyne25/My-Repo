import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Theme Colors (UCU Style)
PURPLE = RGBColor(75, 0, 130)  # Indigo/Purple
YELLOW = RGBColor(255, 215, 0)  # Gold/Yellow
WHITE = RGBColor(255, 255, 255)
DARK_GRAY = RGBColor(50, 50, 50)

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, PURPLE)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    subtitle.text = subtitle_text
    
    # Styling
    title.text_frame.paragraphs[0].font.color.rgb = YELLOW
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.size = Pt(44)
    
    subtitle.text_frame.paragraphs[0].font.color.rgb = WHITE
    subtitle.text_frame.paragraphs[0].font.size = Pt(24)
    
    # Add Logo if exists
    logo_path = 'UCU LOGO.png'
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(0.5), Inches(0.2), height=Inches(0.8))

def add_content_slide(prs, title_text, content_points):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, WHITE)
    
    # Add a top purple bar for theme consistency
    top_bar = slide.shapes.add_shape(
        1,  # Rectangular shape
        Inches(0), Inches(0), prs.slide_width, Inches(1.2)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = PURPLE
    top_bar.line.fill.background()
    
    # Add a bottom yellow bar
    bottom_bar = slide.shapes.add_shape(
        1,
        Inches(0), prs.slide_height - Inches(0.2), prs.slide_width, Inches(0.2)
    )
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = YELLOW
    bottom_bar.line.fill.background()

    # Fixed: Explicitly setting title width and position to prevent vertical text
    title = slide.shapes.title
    title.left = Inches(0.5)
    title.top = Inches(0.2)
    title.width = prs.slide_width - Inches(1.0)
    title.height = Inches(0.8)
    title.text = title_text
    
    title.text_frame.paragraphs[0].font.color.rgb = YELLOW
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.size = Pt(32)
    
    # Fixed: Explicitly setting body width and position
    body_shape = slide.shapes.placeholders[1]
    body_shape.left = Inches(0.5)
    body_shape.top = Inches(1.4)
    body_shape.width = prs.slide_width - Inches(1.0)
    body_shape.height = prs.slide_height - Inches(1.8)
    
    tf = body_shape.text_frame
    tf.word_wrap = True
    
    # Start with an empty text frame
    tf.text = ""
    
    for point in content_points:
        if isinstance(point, list): # Nested points
            for subpoint in point:
                p = tf.add_paragraph()
                p.text = f"• {subpoint}"
                p.level = 1
                p.font.size = Pt(18)
                p.font.color.rgb = DARK_GRAY
        else:
            p = tf.add_paragraph()
            p.text = str(point)
            p.font.size = Pt(20)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(10)

def main():
    prs = Presentation()
    
    # 1. Title Slide
    add_title_slide(prs, 
                   "Departmental Scheduling and Academic Monitoring System\n(DSAMS)", 
                   "A Comprehensive Solution for Academic Operations\nPresentation Report")

    # 2. Introduction
    add_content_slide(prs, "Introduction", [
        "What is DSAMS?",
        "A specialized platform designed to streamline departmental academic monitoring and scheduling.",
        "Targeted at HODs, Lecturers, Coordinators, and Students.",
        "Core Scope:",
        [
            "Timetable automation",
            "Lecture session tracking",
            "Attendance recording & verification",
            "Real-time analytics and alerts"
        ]
    ])

    # 3. Problem Statement
    add_content_slide(prs, "Statement of the Problem", [
        "Challenges at UCU Mbale University College:",
        [
            "Fragmented storage of student lists, timetables and lecturer data.",
            "Manual updates that are time-consuming and error-prone.",
            "Poor communication of schedule changes and academic notices.",
            "Limited visibility and control for HOD over departmental activities.",
            "Over-reliance on informal tools (WhatsApp, printed notices).",
            "Consequences: Missed classes, confusion, and reduced accountability."
        ]
    ])

    # 4. Objectives
    add_content_slide(prs, "Project Objectives", [
        "Main Objective:",
        ["Design and develop DSAMS to automate, centralize and enhance the management of departmental academic operations."],
        "Specific Objectives:",
        [
            "Analyze departmental academic workflows and identify requirements.",
            "Design a role-based system architecture (HOD, Staff, Coordinator, Student).",
            "Implement core modules (Timetable, Monitoring, Attendance, Comms).",
            "Integrate bulk data upload features using standardized templates.",
            "Test and evaluate for usability, accuracy, and efficiency."
        ]
    ])

    # 5. User Roles (RBAC)
    add_content_slide(prs, "User Roles & Permissions", [
        "System Administrator: Full maintenance and global oversight.",
        "Head of Department (HOD): Oversight, Timetabling, Monitoring.",
        "Class Coordinator: Attendance verification & class monitoring.",
        "Lecturer: Schedule viewing & attendance submission.",
        "Student: View-only access to personal and class timetables."
    ])

    # 6. Core Modules
    add_content_slide(prs, "Core Modules", [
        "Scheduling Module: 7-step guided workflow with greedy algorithm for conflict-free timetables.",
        "Hybrid Attendance: Digital capture (web) + hardcopy document verification.",
        "Lecture Monitoring: Real-time status tracking (Conducted, Missed, Pending).",
        "Communication Module: Targeted announcements and automated system alerts."
    ])

    # 7. Key Features
    add_content_slide(prs, "Key Features", [
        "Auto-generation of conflict-free timetables based on capacity.",
        "Premium Glassmorphism UI - responsive for mobile and desktop.",
        "Bulk Data Import: Seamlessly import users/rooms via CSV/Excel.",
        "Automated Compliance Alerts for missed or delayed sessions.",
        "Audit Logging for administrative accountability."
    ])

    # 8. Technology Stack
    add_content_slide(prs, "Technology Stack", [
        "Backend: Python 3.12 / Django 4.2 (LTS)",
        "Frontend: HTML5, Vanilla CSS3 (Custom Glassmorphism), JavaScript (ES6+)",
        "Database: MySQL 8.0+ (Production), SQLite3 (Development)",
        "Libraries: Pandas (Data processing), ReportLab (PDFs), Pillow (Images)",
        "Deployment: Linux/Nginx/Gunicorn ready."
    ])

    # 9. Future Enhancements
    add_content_slide(prs, "Future Roadmap", [
        "Biometric Attendance Integration",
        "AI-Based Student Achievement & Risk Predictions",
        "Mobile Application (Native Android/iOS)",
        "Faculty-wide and Institutional Scalability",
        "Automatic SMS/Email notification gateways"
    ])

    # 10. Conclusion
    add_content_slide(prs, "Conclusion", [
        "DSAMS modernizes academic operations through digital transformation.",
        "Focuses on usability, accountability, and instructional compliance.",
        "Ready for departmental deployment with scalable architecture.",
        "Significant reduction in paper-work and manual intervention."
    ])

    # Save presentation
    output_file = "DSAMS_Presentation.pptx"
    try:
        prs.save(output_file)
        print(f"Presentation saved successfully as {output_file}")
    except PermissionError:
        print(f"ERROR: Permission denied. Please close {output_file} and try again.")

if __name__ == "__main__":
    main()
